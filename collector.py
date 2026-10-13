"""Deterministic public company collector. No AI or API keys."""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
import hashlib
import ipaddress
import json
import logging
from pathlib import Path
import re
import socket
import sqlite3
import subprocess
import time
from urllib.parse import urljoin, urlsplit, urlunsplit, parse_qsl, urlencode
from urllib.robotparser import RobotFileParser
from protego import Protego
import xml.etree.ElementTree as ET

from bs4 import BeautifulSoup
import requests
from openpyxl import load_workbook  # Reading templates/user feedback only.

from app_paths import ROOT, RESOURCES, FROZEN
MSK = timezone(timedelta(hours=3))
LOG = logging.getLogger('collector')
# Use the standard requests client header: Promreg rejects arbitrary custom agents.
UA = requests.utils.default_user_agent()
EMAIL = re.compile(r'[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-zА-Яа-я]{2,}')
PHONE = re.compile(r'(?<!\d)(?:\+7|8)[\s(.-]*\d{3}[\s).\-]*\d{3}[\s.\-]*\d{2}[\s.\-]*\d{2}(?!\d)')


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def text(node):
    return clean(node.get_text(' ', strip=True)) if node else ''


def canonical(url):
    parts = urlsplit(url.strip())
    host = (parts.hostname or '').lower()
    if host in ('www.promreg.ru', 'promreg.ru'):
        host = 'promreg.ru'
    query = [(k, v) for k, v in parse_qsl(parts.query) if not k.lower().startswith('utm_')]
    return urlunsplit(('https' if host == 'promreg.ru' else parts.scheme, host, parts.path or '/', urlencode(sorted(query)), ''))


def website_url(value):
    value = clean(value)
    if not value or ' ' in value or value.startswith(('#', '/')):
        return ''
    return canonical(value if '://' in value else 'https://' + value)


def valid_inn(value):
    value = str(value or '').strip()
    if not value.isdigit() or len(value) not in (10, 12):
        return False
    digits = list(map(int, value))
    def checksum(weights):
        return sum(a*b for a, b in zip(digits, weights)) % 11 % 10
    return (checksum([2,4,10,3,5,9,4,6,8]) == digits[9] if len(value) == 10 else
            checksum([7,2,4,10,3,5,9,4,6,8]) == digits[10] and
            checksum([3,7,2,4,10,3,5,9,4,6,8]) == digits[11])


def contacts(node):
    if node is None:
        return [], []
    body = text(node)
    phones = PHONE.findall(body)
    emails = EMAIL.findall(body)
    for a in node.select('a[href]'):
        href = a['href']
        if href.startswith('tel:'):
            phones += PHONE.findall(href[4:])
            value=href[4:].split('?',1)[0].split(';',1)[0].strip()
            digits=re.sub(r'\D','',value)
            if len(digits)==11 and digits[0] in ('7','8') and re.fullmatch(r'[+\d\s().-]+',value):phones.append(value)
        elif href.startswith('mailto:'):
            emails += EMAIL.findall(href[7:].split('?')[0])
    phones = unique_phones(phones)
    emails = sorted({e.lower() for e in emails if not re.search(r'x{3,}|х{3,}|\*', e, re.I)})
    return phones, emails


def unique_phones(phones):
    result = {}
    for phone in phones:
        digits = re.sub(r'\D', '', phone)[-10:]
        if len(digits) != 10 or len(set(digits)) == 1: continue
        result.setdefault(digits, phone)
    return list(result.values())


def field(soup, selector):
    node = soup.select_one(selector)
    return text(node.select_one('.line, .lines') or node) if node else ''


def labeled_table(node):
    result = {}
    if node:
        for row in node.select('tr'):
            cells = row.find_all(['td', 'th'], recursive=False)
            if len(cells) >= 2:
                result[text(cells[0]).rstrip(':')] = text(cells[1])
    return result


def parse_listing(source, html, url):
    if source['id'] in ('zavody','manufacturers','texzakaz','madeinrussia'):
        from extra_sources import listing
        return listing(source,html,url)
    soup = BeautifulSoup(html, 'html.parser')
    kind = source['id']
    selectors = {'fabricators': 'a.title-relate[href^="/proizvoditel/"]',
                 'productcenter': '.cards .firm .text > a.link[href]',
                 'promreg': '.companybox__title a[href]'}
    companies = []
    for a in soup.select(selectors[kind]):
        href = canonical(urljoin(url, a['href']))
        if kind == 'productcenter' and not re.search(r'/producers/\d+/', href):
            continue
        if href not in companies:
            companies.append(href)
    page_selectors = {'fabricators': '.pager a[href]', 'productcenter': '.paginator a[href], .pagination a[href]',
                      'promreg': 'a.pageline__next2[href], .b-pageline a[href]'}
    pages = []
    seedpath = urlsplit(source['url']).path.rstrip('/')
    for a in soup.select(page_selectors[kind]):
        href = canonical(urljoin(url, a['href']))
        parts = urlsplit(href)
        if parts.hostname != urlsplit(source['url']).hostname:
            continue
        if kind == 'fabricators' and parts.path.rstrip('/') != seedpath:
            continue
        if kind in ('productcenter', 'promreg') and not parts.path.startswith(seedpath + '/page'):
            continue
        if kind == 'productcenter':
            query = dict(parse_qsl(parts.query)); query['sorttype'] = 'createdate'
            href = urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query), ''))
        if href != canonical(url) and href not in pages:
            pages.append(href)
    return companies, pages


def parse_company(source, html, url):
    if source['id'] in ('zavody','manufacturers','texzakaz','madeinrussia'):
        from extra_sources import company
        return company(source,html,url)
    soup = BeautifulSoup(html, 'html.parser')
    kind = source['id']
    item = {'name': '', 'inn': '', 'region': '', 'city': '', 'address': '', 'legal_address': '',
            'status': 'Не проверен', 'feedback': '', 'phones': [], 'emails': [], 'website': '',
            'source': source['name'], 'source_url': canonical(url), 'segment': '', 'categories': [],
            'description': '', 'ogrn': '', 'kpp': '', 'registration_date': '', 'director': '',
            'contact_sources': [], 'status_source': '', 'notes': [], 'extra': {}}
    if kind == 'fabricators':
        item['name'] = field(soup, '.field_jur_lico') or text(soup.select_one('h1'))
        for key, selector in {'inn':'.field_inn','ogrn':'.field_ogrn','kpp':'.field_kpp',
                              'registration_date':'.field_date_register','address':'.field_address'}.items():
            item[key] = field(soup, selector)
        item['region'] = item['address'].split(',')[0].strip()
        item['city'] = item['address'].split(',')[1].strip() if ',' in item['address'] else ''
        phone_node, email_node = soup.select_one('.field_phone'), soup.select_one('.field_email')
        item['phones'] = contacts(phone_node)[0]; item['emails'] = contacts(email_node)[1]
        if phone_node and not item['phones'] or email_node and not item['emails']:
            item['notes'].append('Часть контактов в каталоге скрыта')
        site = soup.select_one('.field_site a[href]')
        item['website'] = website_url(site['href']) if site else ''
        item['description'] = text(soup.select_one('.field_about_company'))
        categories = soup.select('a[href^="/industry/"], a[href^="/produkt/"]')
        item['categories'] = list(dict.fromkeys(text(a) for a in categories if text(a)))
        item['segment'] = 'Туризм и отдых — спальные мешки'
        for node in soup.select('.lines-requisites .fd'):
            label = text(node.select_one('.label'))
            if label: item['extra'][label] = text(node.select_one('.line'))
    elif kind == 'productcenter':
        contact = soup.select_one('.contact_information')
        rows = labeled_table(contact)
        item['name'] = rows.get('Наименование') or text(soup.select_one('h1'))
        for key, label in {'inn':'ИНН','ogrn':'ОГРН','kpp':'КПП','registration_date':'Дата регистрации',
                           'address':'Адрес','legal_address':'Юридический адрес'}.items():
            item[key] = rows.get(label, '')
        item['region'] = text(soup.select_one('.tc_contacts [itemprop="addressRegion"]'))
        item['city'] = text(soup.select_one('.tc_contacts [itemprop="addressLocality"]'))
        item['phones'], item['emails'] = contacts(contact)
        site = soup.select_one('.tc_contacts [itemprop="url"]')
        item['website'] = website_url(text(site)) if site else ''
        item['description'] = text(soup.select_one('.tc_description .iv_text'))
        # Public description can contain extra openly published sales phones.
        p, e = contacts(soup.select_one('.tc_description .iv_text'))
        item['phones'] = unique_phones(item['phones'] + p); item['emails'] = list(dict.fromkeys(item['emails'] + e))
        # Company-linked categories only; never the global site navigation.
        item['extra'] = rows
        # The manufacturer's own product category links provide the segment.
        cats = [a for a in soup.select('.crumb > div > a[href]') if re.match(r'^/producers/catalog-',a['href'])]
        item['categories'] = list(dict.fromkeys(text(a) for a in cats if text(a) and not re.search(r'\d+$', text(a))))
        item['segment'] = item['categories'][0] if item['categories'] else ''
    else:
        node = soup.select_one('.b-catalogbox_detal .catalogbox__content')
        item['name'] = text(node.select_one('h2')) if node else ''
        rows = {}
        for p in node.select('p') if node else []:
            label = text(p.select_one('strong')).rstrip(':')
            if label: rows[label] = clean(text(p)[len(text(p.select_one('strong'))):])
        item['city'] = rows.get('Город', '')
        # Region is the last region breadcrumb/header preceding company title.
        item['region'] = text(soup.select_one('h1')) or item['city']
        item['address'] = rows.get('Адрес', '')
        item['segment'] = rows.get('Вид деятельности', '')
        item['categories'] = [item['segment']] if item['segment'] else []
        item['website'] = website_url(rows.get('Сайт', ''))
        item['phones'], item['emails'] = contacts(node)
        item['description'] = text(soup.select_one('.b-catalogbox_detal .b-editor'))
        item['extra'] = rows
        match = re.search(r'\bИНН\s*[:№]?\s*(\d{10,12})(?!\d)', item['description'])
        if match: item['inn'] = match[1]
    if not item['name'] or len(item['name']) > 500:
        raise ValueError('Не найдено название: разметка карточки изменилась')
    if item['inn'] and not valid_inn(item['inn']):
        item['notes'].append('Опубликованный ИНН не прошёл контрольную сумму: ' + item['inn'])
        item['extra']['Непроверенный ИНН'] = item['inn']; item['inn'] = ''
    if item['phones'] or item['emails']:
        item['contact_sources'].append(canonical(url))
    return item


class Http:
    def __init__(self, config):
        self.config = config
        self.session = requests.Session()
        self.session.headers['User-Agent'] = UA
        self.robots = {}
        self.last = {}
        self.delays={}

    def request(self, url, method='GET', _redirects=0, **kwargs):
        parts = urlsplit(url)
        if parts.scheme not in ('http', 'https') or not parts.hostname:
            raise ValueError('Неподдерживаемый URL')
        host = parts.hostname
        interval=max(self.config['request_delay_seconds'],self.delays.get(host,0))
        delay = interval - (time.monotonic() - self.last.get(host, 0))
        if delay > 0: time.sleep(delay)
        try:
            r = self.session.request(method, url, timeout=self.config['timeout_seconds'], allow_redirects=False, **kwargs)
        finally:
            self.last[host] = time.monotonic()
        if r.status_code in (301,302,303,307,308):
            target = urljoin(url, r.headers['Location'])
            if urlsplit(target).scheme == 'http' and parts.scheme == 'https':
                target = 'https:' + target[5:]
            if _redirects >= 5:
                raise ValueError('Слишком много перенаправлений')
            if not public_website(target):
                raise ValueError('Перенаправление не ведёт на публичный сайт')
            return self.request(target, method='GET', _redirects=_redirects+1, **kwargs)
        r.raise_for_status()
        r.encoding = 'utf-8' if b'charset=utf-8' in r.content[:3000].lower() or b'charset="utf-8"' in r.content[:3000].lower() else r.apparent_encoding
        return r

    def allowed(self, url):
        parts = urlsplit(url); root = parts.scheme + '://' + parts.netloc
        if root not in self.robots:
            try:
                r = self.request(root + '/robots.txt')
                rp=Protego.parse(r.text)
            except requests.HTTPError as ex:
                if ex.response.status_code == 404: rp=Protego.parse('')
                else: raise
            self.robots[root] = rp
            interval=rp.crawl_delay(UA) or 0
            rate=rp.request_rate(UA)
            if rate and rate.requests:interval=max(interval,rate.seconds/rate.requests)
            self.delays[parts.hostname]=max(self.delays.get(parts.hostname,0),interval)
        return self.robots[root].can_fetch(url,UA)

    def get(self, url):
        if not self.allowed(url):
            raise ValueError('Адрес запрещён robots.txt: ' + url)
        response=self.request(url)
        host=urlsplit(url).hostname
        # This public host serves a fixed browser-cookie preface, not a CAPTCHA.
        # Honor that known literal cookie without evaluating third-party JavaScript.
        if host=='xn--80aegj1b5e.xn--p1ai' and "document.cookie='beget=begetok'" in response.text and 'location.reload()' in response.text:
            self.session.cookies.set('beget','begetok',domain=host,path='/')
            response=self.request(url)
        return response


def public_website(url):
    parts = urlsplit(url)
    if parts.scheme not in ('http','https') or not parts.hostname or parts.username or parts.port not in (None,80,443):
        return False
    try:
        return all(ipaddress.ip_address(record[4][0]).is_global for record in socket.getaddrinfo(parts.hostname, None))
    except (OSError, ValueError):
        return False


def enrich_contacts(item, http):
    root = item['website']
    if not root or not public_website(root): return
    pages = [root]; visited = set()
    for _ in range(3):
        if not pages: break
        url = pages.pop(0)
        if url in visited: continue
        visited.add(url)
        try:
            response = http.get(url)
            soup = BeautifulSoup(response.text, 'html.parser')
            for t in soup.select('script, style, noscript'): t.decompose()
            phones, emails = contacts(soup)
            if phones or emails:
                item['phones'] = unique_phones(item['phones'] + phones)
                item['emails'] = list(dict.fromkeys(item['emails'] + emails))
                item['contact_sources'].append(url)
            # Enrich missing INN only from explicitly labeled official requisites.
            if not item['inn']:
                candidates = re.findall(r'\bИНН\s*[:№]?\s*(\d{10,12})(?!\d)', text(soup), re.I)
                candidates = list(dict.fromkeys(i for i in candidates if valid_inn(i)))
                if len(candidates) == 1:
                    item['inn'] = candidates[0]
                    item['extra']['Источник ИНН'] = url
                elif len(candidates) > 1:
                    item['notes'].append('На официальном сайте несколько ИНН; нужен выбор юрлица')
            for a in soup.select('a[href]'):
                href = canonical(urljoin(url, a['href']))
                if urlsplit(href).hostname != urlsplit(root).hostname: continue
                if re.search(r'контакт|реквизит|contact|requisit', text(a)+' '+href, re.I) and href not in visited and href not in pages:
                    pages.append(href)
        except Exception as ex:
            LOG.info('Официальный сайт %s: %s', url, ex)
            item['notes'].append('Не удалось открыть страницу официального сайта: ' + url)
    item['contact_sources'] = list(dict.fromkeys(item['contact_sources']))


def classify_fns(row):
    if row.get('v'): return 'Регистрация признана недействительной', 'Дата: ' + row['v']
    if row.get('e'): return 'Недействующая', 'Дата прекращения деятельности: ' + row['e']
    if row.get('k') not in ('ul','fl') or not row.get('o') or not row.get('r'):
        return 'Не проверен', 'Неполный ответ ФНС'
    return 'Действующая', 'В поиске ФНС нет даты прекращения деятельности. Процедуры ликвидации, банкротства и реорганизации требуют проверки выписки.'


def check_fns(item, http):
    if not valid_inn(item['inn']):
        item['notes'].append('Проверка статуса невозможна без подтверждённого ИНН')
        return
    base = 'https://egrul.nalog.ru/'
    try:
        http.request(base + 'index.html')
        result = http.request(base, method='POST', data={'query':item['inn'], 'region':'', 'PreventChromeAutocomplete':''}).json()
        if result.get('captchaRequired') or not result.get('t'):
            raise ValueError('ФНС запросила капчу либо не вернула идентификатор поиска')
        answer = None
        for _ in range(4):
            answer = http.request(base + 'search-result/' + result['t']).json()
            if answer.get('status') != 'wait': break
            time.sleep(1)
        rows = [row for row in (answer or {}).get('rows', []) if row.get('i') == item['inn']]
        if item['ogrn']: rows = [row for row in rows if row.get('o') == item['ogrn']]
        if len(rows) != 1:
            raise ValueError('Нет однозначного совпадения по ИНН и ОГРН')
        row = rows[0]
        item['status'], note = classify_fns(row)
        item['notes'].append(note)
        item['status_source'] = base + 'index.html — ' + datetime.now(MSK).isoformat(timespec='seconds')
        item['extra']['ФНС'] = row
        for key, fieldname in {'ogrn':'o','kpp':'p','registration_date':'r','director':'g'}.items():
            if row.get(fieldname): item[key] = row[fieldname]
        if row.get('rn'): item['region'] = row['rn']
    except Exception as ex:
        item['notes'].append('ФНС: статус не проверен — ' + str(ex))
        LOG.warning('ФНС, ИНН %s: %s', item['inn'], ex)


def normalized_name(name):
    return re.sub(r'[^\w]', '', name.lower().replace('ё','е'))


def identities(item):
    keys = ['url:' + item['source_url']] if item.get('source_url') else []
    if valid_inn(item.get('inn')): keys.append('inn:' + item['inn'])
    name = normalized_name(item['name'])
    region = clean(item.get('region')).lower()
    # Weak keys are only used for companies without verified INN.
    if not item.get('inn'):
        if region: keys.append('name:' + name + ':' + region)
        if item.get('website'): keys.append('site:' + urlsplit(item['website']).hostname + ':' + name)
    return keys


def signatures(item):
    name = normalized_name(item['name'])
    region = re.sub(r'[^\w]', '', clean(item.get('region')).lower())
    region = re.sub(r'^г(?:ород)?(?=москва|санкт)', '', region)
    result = ['name-region:'+name+':'+region] if region else []
    if item.get('website'):
        host=(urlsplit(item['website']).hostname or '').removeprefix('www.').encode('idna').decode('ascii')
        result.append('name-site:'+name+':'+host)
    return result


class CallConflict(ValueError):
    pass


def call_fields(row):
    return dict(called=bool(row['called']) if row else False,
                call_note=row['call_note'] if row else '',called_at=row['called_at'] if row else '',
                call_updated_at=row['updated_at'] if row else '',call_revision=row['revision'] if row else 0)


class Store:
    def __init__(self, path):
        self.db = sqlite3.connect(path)
        self.db.row_factory = sqlite3.Row
        self.db.executescript('''
          PRAGMA journal_mode=WAL;
          CREATE TABLE IF NOT EXISTS companies(id INTEGER PRIMARY KEY, added_day TEXT, added_at TEXT, payload TEXT NOT NULL);
          CREATE TABLE IF NOT EXISTS identities(key TEXT PRIMARY KEY, company_id INTEGER NOT NULL);
          CREATE TABLE IF NOT EXISTS candidates(url TEXT PRIMARY KEY, source TEXT, state TEXT DEFAULT 'pending', attempts INTEGER DEFAULT 0, last_attempt TEXT, error TEXT);
          CREATE TABLE IF NOT EXISTS listing_pages(url TEXT PRIMARY KEY, source TEXT, last_scan TEXT);
          CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY, started TEXT, finished TEXT, added INTEGER, errors TEXT);
          CREATE TABLE IF NOT EXISTS signatures(key TEXT, company_id INTEGER, PRIMARY KEY(key,company_id));
          CREATE INDEX IF NOT EXISTS signature_key ON signatures(key);
          CREATE TABLE IF NOT EXISTS company_calls(company_id INTEGER PRIMARY KEY REFERENCES companies(id),
            called INTEGER NOT NULL,call_note TEXT NOT NULL,called_at TEXT NOT NULL,updated_at TEXT NOT NULL,revision INTEGER NOT NULL);
        ''')
        if 'payload' not in {row[1] for row in self.db.execute('PRAGMA table_info(candidates)')}:
            self.db.execute('ALTER TABLE candidates ADD COLUMN payload TEXT');self.db.commit()
        with self.db:
            for row in self.db.execute('SELECT id,payload FROM companies WHERE id NOT IN (SELECT company_id FROM signatures)'):
                for key in signatures(json.loads(row['payload'])):
                    self.db.execute('INSERT OR IGNORE INTO signatures VALUES (?,?)',(key,row['id']))

    def today_count(self, day):
        return self.db.execute('SELECT COUNT(*) FROM companies WHERE added_day=?', (day,)).fetchone()[0]

    def add(self, item, day=None, limit=None):
        keys = identities(item)
        with self.db:
            existing = self.db.execute('SELECT company_id FROM identities WHERE key IN ('+','.join('?' for _ in keys)+')', keys).fetchone() if keys else None
            if not existing:
                possible = set()
                for key in signatures(item):
                    for row in self.db.execute('SELECT c.id,c.payload FROM signatures s JOIN companies c ON c.id=s.company_id WHERE s.key=?',(key,)):
                        prior=json.loads(row['payload'])
                        if valid_inn(item.get('inn')) and valid_inn(prior.get('inn')) and item['inn']!=prior['inn']: continue
                        possible.add(row['id'])
                if len(possible)==1: existing=(possible.pop(),)
            if existing:
                # Remember a duplicate source URL so it is not re-collected tomorrow.
                for key in keys: self.db.execute('INSERT OR IGNORE INTO identities VALUES (?,?)',(key,existing[0]))
                prior=json.loads(self.db.execute('SELECT payload FROM companies WHERE id=?',(existing[0],)).fetchone()[0])
                for fieldname in ('inn','ogrn','kpp','registration_date','director','address','legal_address','description','website','segment'):
                    if not prior.get(fieldname) and item.get(fieldname): prior[fieldname]=item[fieldname]
                if item.get('status')!='Не проверен' and item.get('status_source'):
                    prior['status']=item['status'];prior['status_source']=item['status_source']
                prior['phones']=unique_phones(prior.get('phones',[])+item.get('phones',[]))
                for fieldname in ('emails','contact_sources','categories','notes'):
                    prior[fieldname]=list(dict.fromkeys(prior.get(fieldname,[])+item.get(fieldname,[])))
                prior.setdefault('extra',{}).setdefault('Другие источники',[])
                provenance={'name':item.get('source',''),'url':item.get('source_url','')}
                if provenance['url'] and provenance not in prior['extra']['Другие источники']:
                    prior['extra']['Другие источники'].append(provenance)
                for key in signatures(item)+signatures(prior):
                    self.db.execute('INSERT OR IGNORE INTO signatures VALUES (?,?)',(key,existing[0]))
                self.db.execute('UPDATE companies SET payload=? WHERE id=?',(json.dumps(prior,ensure_ascii=False),existing[0]))
                return False
            if day and limit is not None and self.today_count(day) >= limit: return False
            stamp = datetime.now(MSK).isoformat(timespec='seconds') if day else ''
            item['added_at'] = stamp
            cur = self.db.execute('INSERT INTO companies(added_day,added_at,payload) VALUES (?,?,?)',(day,stamp,json.dumps(item,ensure_ascii=False)))
            for key in keys: self.db.execute('INSERT INTO identities VALUES (?,?)',(key,cur.lastrowid))
            for key in signatures(item): self.db.execute('INSERT OR IGNORE INTO signatures VALUES (?,?)',(key,cur.lastrowid))
        return True

    def items(self):
        calls={row['company_id']:row for row in self.db.execute('SELECT * FROM company_calls')}
        return [dict(json.loads(row['payload']), id=row['id'],**call_fields(calls.get(row['id']))) for row in self.db.execute('SELECT * FROM companies ORDER BY id')]

    def update_call(self,company_id,called,note,revision):
        # Separate employee data from scraped payloads. Lock before checking a revision.
        with self.db:
            self.db.execute('BEGIN IMMEDIATE')
            if not self.db.execute('SELECT 1 FROM companies WHERE id=?',(company_id,)).fetchone():raise KeyError(company_id)
            old=call_fields(self.db.execute('SELECT * FROM company_calls WHERE company_id=?',(company_id,)).fetchone())
            if revision!=old['call_revision']:raise CallConflict('Запись уже изменена в другом окне. Откройте карточку заново; ваш текст остаётся в поле.')
            if called==old['called'] and note==old['call_note']:return old
            now=datetime.now(MSK).isoformat()
            called_at=(old['called_at'] or now) if called else ''
            self.db.execute('INSERT INTO company_calls VALUES (?,?,?,?,?,?) ON CONFLICT(company_id) DO UPDATE SET called=excluded.called,call_note=excluded.call_note,called_at=excluded.called_at,updated_at=excluded.updated_at,revision=excluded.revision',
                            (company_id,int(called),note,called_at,now,revision+1))
        return dict(called=called,call_note=note,called_at=called_at,call_updated_at=now,call_revision=revision+1)


def import_template(store, path):
    if store.db.execute('SELECT COUNT(*) FROM companies').fetchone()[0]: return
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    for row in list(sheet.values)[1:]:
        if len(row) < 2 or not row[1]: continue
        values = list(row) + [None] * 12
        item = {'name':clean(values[1]), 'inn':clean(values[2]), 'region':clean(values[3]),
                'status':'Не проверен', 'feedback':clean(values[0]), 'phones':str(values[5] or '').splitlines(),
                'emails':str(values[6] or '').splitlines(), 'website':website_url(values[7]),
                'source':clean(values[8]), 'source_url':'', 'segment':clean(values[9]),
                'notes':['Исходная таблица; статус из неё не проверен: '+clean(values[4])],
                'extra':{'Примечание исходной таблицы':clean(values[10])}}
        store.add(item)
    wb.close()


def feedback_from_excel(store, output):
    if not output.exists(): return
    wb = load_workbook(output, read_only=True, data_only=True)
    rows = iter(wb.active.values); headers = next(rows)
    indexes = {str(v):i for i,v in enumerate(headers) if v}
    if 'Обратная связь' not in indexes: wb.close(); return
    items = store.items()
    by_inn = {v['inn']:v for v in items if v.get('inn')}
    by_name = {v['name']:v for v in items}
    with store.db:
        for row in rows:
            inn = clean(row[indexes['ИНН']]); name = clean(row[indexes['Наименование']])
            item = by_inn.get(inn) if inn else by_name.get(name)
            if item:
                item['feedback'] = clean(row[indexes['Обратная связь']])
                store.db.execute('UPDATE companies SET payload=? WHERE id=?',(json.dumps(item,ensure_ascii=False),item['id']))
    wb.close()


def export(store, config):
    output = ROOT / config['output']; output.parent.mkdir(parents=True, exist_ok=True)
    feedback_file=output
    manifest=Path(str(output)+'.latest.json')
    if manifest.exists():
        name=json.loads(manifest.read_text(encoding='utf-8')).get('file','')
        if name and Path(name).name==name and name.endswith('.xlsx'):
            feedback_file=output.parent/name
    feedback_from_excel(store, feedback_file)
    payload = output.parent / 'companies.json'
    temp = payload.with_suffix('.tmp')
    temp.write_text(json.dumps(store.items(),ensure_ascii=False,indent=2),encoding='utf-8'); temp.replace(payload)
    if FROZEN or config.get('export_engine')=='python':
        from excel_export import write_workbook
        actual=write_workbook(store.items(),output,RESOURCES/'timezones.json')
        LOG.info('Excel: %d компаний; %s',len(store.items()),actual)
        return
    result = subprocess.run([config['node'],str(ROOT/'export.mjs'),str(payload),str(output)],cwd=ROOT,capture_output=True,text=True,encoding='utf-8')
    if result.returncode: raise RuntimeError('Ошибка экспорта Excel: '+result.stderr[-2500:])
    LOG.info(result.stdout.strip())


@contextmanager
def single_run(path):
    path.parent.mkdir(parents=True,exist_ok=True)
    handle = path.open('a+b')
    try:
        if handle.tell() == 0: handle.write(b'0'); handle.flush()
        handle.seek(0)
        if __import__('os').name == 'nt':
            import msvcrt
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        handle.close(); raise RuntimeError('Сборщик уже запущен')
    try: yield
    finally:
        handle.seek(0)
        if __import__('os').name == 'nt': msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        else: fcntl.flock(handle, fcntl.LOCK_UN)
        handle.close()


def iter_discover(store, source, http, max_pages=None):
    if source['id']=='egrul':
        from fns_registry import discover as registry_discover
        yield from registry_discover(store,source,http,max_pages)
        return
    seed = canonical(source['url'])
    with store.db:
        store.db.execute('INSERT OR IGNORE INTO listing_pages(url,source) VALUES (?,?)',(seed,source['id']))
    # Rescan first page every day; rotate other pages to eventually cover the catalogue.
    pages = [seed] + [r[0] for r in store.db.execute('SELECT url FROM listing_pages WHERE source=? AND url<>? ORDER BY last_scan IS NOT NULL,last_scan,url',(source['id'],seed))]
    pages.extend(page for page in source.get('sitemaps',[]) if page not in pages)
    visited = set()
    while pages and (max_pages is None or len(visited)<max_pages):
        url = pages.pop(0)
        if url in visited: continue
        visited.add(url)
        if hasattr(http,'allowed') and not http.allowed(url):
            LOG.info('%s: robots.txt ограничивает страницу %s',source['name'],url)
            continue
        response = http.get(url)
        if urlsplit(url).path.endswith('.xml'):
            from extra_sources import card_url
            xml=ET.fromstring(response.content)
            locs=[canonical(node.text.strip()) for node in xml.iter() if node.tag.rsplit('}',1)[-1]=='loc' and node.text]
            locs=[loc for loc in locs if urlsplit(loc).hostname==urlsplit(seed).hostname]
            if xml.tag.rsplit('}',1)[-1]=='sitemapindex':companies=[];next_pages=locs
            else:companies=[loc for loc in locs if card_url(source,loc)];next_pages=[]
        else:
            companies, next_pages = parse_listing(source,response.text,url)
            if not companies: raise ValueError('На странице каталога не найдены карточки: '+url)
        with store.db:
            for company in companies:
                store.db.execute('INSERT OR IGNORE INTO candidates(url,source) VALUES (?,?)',(company,source['id']))
            for page in next_pages:
                store.db.execute('INSERT OR IGNORE INTO listing_pages(url,source) VALUES (?,?)',(page,source['id']))
                if page not in pages and page not in visited: pages.append(page)
            store.db.execute('UPDATE listing_pages SET last_scan=? WHERE url=?',(datetime.now(MSK).isoformat(),url))
        LOG.info('%s: %s, найдено %d карточек',source['name'],url,len(companies))
        yield len(companies)


def discover(store, source, http, max_pages):
    for _ in iter_discover(store,source,http,max_pages): pass


def run(config, export_only=False, batch_size=10):
    data = ROOT/'data'; data.mkdir(exist_ok=True)
    with single_run(data/'run.lock'):
        store = Store(data/'companies.sqlite3')
        if config.get('template'):import_template(store, Path(config['template']))
        if export_only:
            export(store,config); return 0
        started = datetime.now(MSK).isoformat(timespec='seconds'); day = started[:10]
        limit = config['daily_limit']; before = store.today_count(day); before_total=len(store.items()); errors = []
        http = Http(config)
        if limit is not None and before >= limit:
            LOG.info('Сегодня уже добавлено %d компаний; дневной лимит %d',before,limit)
            export(store,config); return 0
        attempts = 0
        exhausted = set()
        discovery={source['id']:iter_discover(store,source,http,config['max_listing_pages_per_source']) for source in config['sources']}
        max_attempts=config['max_company_attempts_per_run']
        def batch_done():
            return batch_size is not None and store.db.execute('SELECT COUNT(*) FROM companies').fetchone()[0]-before_total>=batch_size
        def next_candidate(source):
            return store.db.execute("SELECT * FROM candidates WHERE source=? AND state IN ('pending','failed') AND (last_attempt IS NULL OR last_attempt<>?) ORDER BY attempts,rowid LIMIT 1",(source['id'],started)).fetchone()
        while not batch_done() and (limit is None or store.today_count(day)<limit) and (max_attempts is None or attempts<max_attempts):
            progressed = False
            for source in config['sources']:
                if source['id'] in exhausted: continue
                if batch_done() or (limit is not None and store.today_count(day)>=limit) or (max_attempts is not None and attempts>=max_attempts): break
                # Each run retries prior failures once. Successful cards are never re-added.
                candidate=next_candidate(source)
                if not candidate:
                    try:
                        next(discovery[source['id']])
                        progressed=True
                        candidate=next_candidate(source)
                    except StopIteration:
                        exhausted.add(source['id']);continue
                    except Exception as ex:
                        errors.append(source['name']+': '+str(ex));LOG.warning('%s: %s',source['name'],ex)
                        exhausted.add(source['id']);continue
                if not candidate: continue
                progressed=True; attempts+=1; url=candidate['url']
                try:
                    if source['id']=='egrul':
                        item=json.loads(candidate['payload'])
                        if not item.get('inn') or not valid_inn(item['inn']):raise ValueError('Нет подтверждённого ИНН в ответе ЕГРЮЛ')
                    else:
                        response = http.get(url)
                        archive = data/'pages'; archive.mkdir(exist_ok=True)
                        filename = hashlib.sha256(url.encode()).hexdigest()[:20]+'.html'
                        (archive/filename).write_text(response.text,encoding='utf-8')
                        item = parse_company(source,response.text,url)
                        item['extra']['Архив карточки'] = filename
                        if source['id']=='zavody':
                            details_url=url.rstrip('/')+'/details'
                            try:
                                details=http.get(details_url)
                                more=parse_company(source,details.text,url)
                                for key in ('name','inn','ogrn','kpp'):
                                    if more.get(key):item[key]=more[key]
                                item['extra'].update(more['extra'])
                                item['extra']['Источник реквизитов']=details_url
                                (archive/(filename.removesuffix('.html')+'-details.html')).write_text(details.text,encoding='utf-8')
                            except Exception as ex:item['notes'].append('Реквизиты каталога недоступны: '+str(ex))
                    LOG.info('Обработка: %s',item['name'])
                    if config['enrich_official_sites']: enrich_contacts(item,http)
                    if config['check_fns'] and source['id']!='egrul': check_fns(item,http)
                    added = store.add(item,datetime.now(MSK).date().isoformat(),limit)
                    with store.db:
                        store.db.execute("UPDATE candidates SET state=?,attempts=attempts+1,last_attempt=?,error=NULL WHERE url=?",('done' if added else 'duplicate',started,url))
                    LOG.info('%s: %s | %s | телефон %d, почта %d', 'Добавлено' if added else 'Повтор',item['name'],item['status'],len(item['phones']),len(item['emails']))
                except Exception as ex:
                    errors.append(url+': '+str(ex)); LOG.warning('Карточка %s: %s',url,ex)
                    with store.db:
                        store.db.execute("UPDATE candidates SET state='failed',attempts=attempts+1,last_attempt=?,error=? WHERE url=?",(started,str(ex),url))
                    if isinstance(ex,requests.HTTPError) and ex.response.status_code in (403,429): exhausted.add(source['id'])
            if not progressed: break
        added = len(store.items())-before_total
        with store.db:
            store.db.execute('INSERT INTO runs(started,finished,added,errors) VALUES (?,?,?,?)',(started,datetime.now(MSK).isoformat(),added,json.dumps(errors,ensure_ascii=False)))
        export(store,config)
        LOG.info('За запуск добавлено: %d. Всего сегодня: %d. Ошибок: %d.',added,store.today_count(day),len(errors))
        if batch_size is not None and added<batch_size:
            LOG.warning('Запрошено %d новых компаний, добавлено %d. Очередь исчерпана или источники недоступны.',batch_size,added)
            return 2
        return 0 if batch_size is not None or not errors else 2


def main():
    parser=argparse.ArgumentParser(description='Сбор всех новых компаний без ИИ и повторов')
    parser.add_argument('--config',type=Path,default=ROOT/'config.json')
    parser.add_argument('--export-only',action='store_true')
    parser.add_argument('--limit',type=int,default=10,help='Количество новых компаний за запуск, без дневного ограничения')
    parser.add_argument('--all',action='store_true',help='Продолжать до исчерпания всех разделов')
    parser.add_argument('--check-sources',action='store_true',help='Проверить каталоги и по одной карточке, без добавления компаний')
    args=parser.parse_args()
    (ROOT/'logs').mkdir(exist_ok=True)
    logging.basicConfig(level=logging.INFO,format='%(asctime)s %(levelname)s %(message)s',handlers=[logging.FileHandler(ROOT/'logs'/'collector.log',encoding='utf-8'),logging.StreamHandler()])
    try:
        config=json.loads(args.config.read_text(encoding='utf-8'))
        if config['daily_limit'] is not None and (not isinstance(config['daily_limit'],int) or config['daily_limit']<1): raise ValueError('daily_limit должен быть положительным числом или null')
        if args.check_sources:
            http=Http(config); failures=0
            for source in config['sources']:
                try:
                    if source['id']=='egrul':
                        from fns_registry import search,company as registry_company
                        http.request('https://egrul.nalog.ru/index.html')
                        rows=search(http,source.get('queries',['завод'])[0],source.get('region',''))
                        examples=[row for row in rows if row.get('k')=='ul']
                        if not examples:raise ValueError('Нет юридических лиц в результатах ЕГРЮЛ')
                        item=registry_company(examples[0],source)
                        LOG.info('Проверен %s: %d карточек, пример %s',source['name'],len(examples),item['name'])
                        continue
                    page=http.get(source['url'])
                    urls,pages=parse_listing(source,page.text,source['url'])
                    if not urls: raise ValueError('Нет карточек в каталоге')
                    item=parse_company(source,http.get(urls[0]).text,urls[0])
                    LOG.info('Проверен %s: %d карточек, пагинация %d, пример %s, сегмент %s',source['name'],len(urls),len(pages),item['name'],item['segment'])
                except Exception as ex:
                    failures+=1;LOG.error('%s: %s',source['name'],ex)
            return 1 if failures else 0
        if args.limit<1: raise ValueError('--limit должен быть положительным')
        return run(config,args.export_only,None if args.all else args.limit)
    except Exception:
        LOG.exception('Сбор завершился с ошибкой. Уже собранные данные остаются в SQLite.'); return 1


if __name__ == '__main__':
    raise SystemExit(main())
