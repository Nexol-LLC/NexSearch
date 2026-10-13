from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import sys
import os
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
import pytest
from fastapi.testclient import TestClient
from psycopg.errors import ForeignKeyViolation
from network.postgres_runtime import start_postgres,stop_postgres
from network.database import Database,Conflict
from network.api import create_app
from network.auth import Authentication,HASHER

ROOT=Path(__file__).resolve().parents[2]

@pytest.fixture(scope='session')
def pg():
    home=(ROOT/'network/.runtime/test-host').resolve()
    assert home.is_relative_to((ROOT/'network/.runtime').resolve())
    runtime=Path(os.environ.get('NEXSEARCH_TEST_POSTGRES',ROOT/'network/.runtime/postgres'))
    dsn=start_postgres(home,runtime,int(os.environ.get('NEXSEARCH_TEST_PORT','54354')))
    database=Database(dsn);database.initialize()
    with database.transaction(admin=True) as conn:conn.execute((ROOT/'network/schema14.sql').read_text(encoding='utf-8'))
    yield database
    database.close()
    if (home/'postgres-data/postmaster.pid').exists():stop_postgres(home,runtime)

@pytest.fixture
def db(pg):
    with pg.transaction(admin=True) as conn:
        conn.execute('TRUNCATE employees,companies,jobs,carriers,scraper_candidates,scraper_pages,audit_events,company_checks CASCADE')
    yield pg

def item(n,inn=''):
    from datetime import datetime,timezone
    if not inn:
        digits='7700'+str(n).zfill(5);inn=digits+str(sum(int(d)*w for d,w in zip(digits,[2,4,10,3,5,9,4,6,8]))%11%10)
    return dict(name='Компания '+str(n),inn=inn,region='Москва',phones=[],emails=[],notes=[],extra={},
        website='',source='Каталог',source_url='https://example.org/company/'+str(n),status='Действует',contact_sources=[],
        verification={'active':True,'okved':'43.21','checked_at':datetime.now(timezone.utc).isoformat()},
        financial={'year':2025,'revenue_rub':50_000_000,'reports':{'2025':{'2110':50_000_000}}})

def users(db):
    auth=Authentication(db)
    alice=auth.register('Анна','Первая','password-one','device-a','Computer A','127.0.0.1')
    bob=auth.register('Борис','Второй','password-two','device-b','Computer B','127.0.0.1')
    with db.transaction(admin=True) as conn:conn.execute('UPDATE employees SET approved=true')
    return alice,bob

def http(db):
    return TestClient(create_app(db,start_worker=False,pack=ROOT/'pack'))
def header(person):return {'Authorization':'Bearer '+person['token']}

def test_postgres_role_and_password_hashes_are_real_and_not_returned(db):
    alice,bob=users(db)
    assert 'password_hash' not in alice['employee'] and 'password' not in alice
    with db.transaction(admin=True) as conn:
        assert conn.execute('SELECT rolsuper FROM pg_roles WHERE rolname=current_user').fetchone()['rolsuper'] is False
        hashes=conn.execute('SELECT password_hash FROM employees ORDER BY created_at').fetchall()
    assert all(r['password_hash'].startswith('$argon2id$') for r in hashes)
    assert HASHER.verify(hashes[0]['password_hash'],'password-one')
    app=http(db)
    assert app.get('/api/me',headers=header(alice)).json()['employee']['login']=='pervaya.anna'
    assert app.get('/api/me').status_code==401
    wrong=app.post('/api/auth/login',json=dict(login='pervaya.anna',password='wrong-password',device_id='d'))
    assert wrong.status_code==401

def test_concurrent_claims_are_disjoint_and_rls_filters_private_data(db):
    alice,bob=users(db);a=alice['employee']['id'];b=bob['employee']['id']
    for n in range(20):db.ingest(item(n))
    with ThreadPoolExecutor(max_workers=2) as pool:
        one,two=list(pool.map(lambda employee:db.claim_available(employee,10),[a,b]))
    assert len(one)==len(two)==10 and not set(one)&set(two)
    assert len(db.leads(a))==len(db.leads(b))==10
    with db.transaction(a) as conn:
        assert conn.execute('SELECT count(*) AS n FROM employee_leads').fetchone()['n']==10
        assert all(str(r['employee_id'])==str(a) for r in conn.execute('SELECT employee_id FROM employee_leads'))
    api=http(db)
    denied=api.patch('/api/leads/'+two[0],headers=header(alice),json=dict(status='taken',called=True,call_note='stolen',revision=0))
    assert denied.status_code==404
    assert api.get('/api/admin/employees',headers=header(alice)).status_code==403

def test_take_requires_call_and_creates_only_owners_client_with_free_stage(db):
    alice,bob=users(db);a=alice['employee']['id'];b=bob['employee']['id']
    company,_,_=db.ingest(item(1),a)
    app=http(db)
    assert app.patch('/api/leads/'+str(company),headers=header(alice),json=dict(status='taken',called=False,call_note='',revision=0)).status_code==400
    result=app.patch('/api/leads/'+str(company),headers=header(alice),json=dict(status='taken',called=True,call_note='Согласовали расчёт',revision=0))
    assert result.status_code==200 and result.json()['called_at']
    mine=app.get('/api/clients',headers=header(alice)).json()['companies']
    assert len(mine)==1 and not app.get('/api/clients',headers=header(bob)).json()['companies']
    stage='Ждём документы / согласование 2'
    assert app.patch('/api/clients/'+str(company)+'/stage',headers=header(alice),json=dict(stage=stage,revision=0)).status_code==200
    assert app.get('/api/clients?order=stage',headers=header(alice)).json()['companies'][0]['stage']==stage
    assert stage not in app.get('/api/clients/stages',headers=header(bob)).json()['stages']
    assert app.patch('/api/clients/'+str(company)+'/stage',headers=header(bob),json=dict(stage='foreign',revision=1)).status_code==404
    assert app.patch('/api/clients/'+str(company)+'/stage',headers=header(alice),json=dict(stage='stale',revision=0)).status_code==409
    app.patch('/api/leads/'+str(company),headers=header(alice),json=dict(status='refused',called=True,call_note='Отказ',revision=1))
    assert app.get('/api/clients',headers=header(alice)).json()['companies']==[]
    assert db.claim_available(b,10)==[]  # Refusal never reassigns the lead to a different person.
    app.patch('/api/leads/'+str(company),headers=header(alice),json=dict(status='taken',called=True,call_note='Передумали',revision=2))
    assert app.get('/api/clients',headers=header(alice)).json()['companies'][0]['stage']==stage

def test_carriers_are_shared_and_updates_detect_conflicts(db):
    alice,bob=users(db);app=http(db)
    payload=dict(name='Транспорт',contacts='+7 900 123-45-67',routes='Москва — Казань\nКазань — Самара',description='Тент и трал')
    carrier=app.post('/api/carriers',headers=header(alice),json=payload).json()
    assert app.get('/api/carriers',headers=header(bob)).json()['carriers'][0]['routes']==payload['routes']
    updated=app.patch('/api/carriers/'+carrier['id'],headers=header(bob),json=dict(payload,description='Общий комментарий',revision=0))
    assert updated.status_code==200
    assert app.patch('/api/carriers/'+carrier['id'],headers=header(alice),json=dict(payload,revision=0)).status_code==409

def test_deduplication_never_overwrites_call_or_assigns_to_second_employee(db):
    alice,bob=users(db);a=alice['employee']['id'];b=bob['employee']['id']
    company,_,assigned=db.ingest(item(1,'5402176081'),a)
    assert assigned
    db.update_lead(a,company,'taken',True,'В работе у Анны',0)
    enriched=dict(item(2,'5402176081'),phones=['+7 495 123-45-67'],name='Другое имя в каталоге')
    duplicate,new,assigned=db.ingest(enriched,b)
    assert duplicate==company and not new and not assigned
    assert db.leads(b)==[]
    assert db.leads(a,clients=True)[0]['call_note']=='В работе у Анны'
    assert db.leads(a)[0]['phones']==enriched['phones']

def test_private_plans_and_export_do_not_leak_between_employees(db):
    alice,bob=users(db);app=http(db)
    plan=app.post('/api/plans',headers=header(alice),json=dict(name='Погрузка Анны',project={'boxes':[{'name':'cargo'}]})).json()
    assert app.get('/api/plans/'+plan['id'],headers=header(bob)).status_code==404
    assert app.get('/api/plans',headers=header(bob)).json()['plans']==[]
    company,_,_=db.ingest(item(1),alice['employee']['id']);db.update_lead(alice['employee']['id'],company,'taken',True,'only Alice',0)
    import io
    from openpyxl import load_workbook
    for person,expected in [(alice,2),(bob,1)]:
        download=app.get('/api/export/clients',headers=header(person))
        wb=load_workbook(io.BytesIO(download.content));assert wb.active.max_row==expected;wb.close()

def test_admin_sees_employee_bases_and_can_revoke_access(db):
    alice,bob=users(db);app=http(db)
    with db.transaction(admin=True) as conn:conn.execute('UPDATE employees SET is_admin=true WHERE id=%s',(alice['employee']['id'],))
    company,_,_=db.ingest(item(1),bob['employee']['id']);db.update_lead(bob['employee']['id'],company,'taken',True,'Bob client',0)
    assert app.get('/api/admin/employees',headers=header(alice)).status_code==200
    assert len(app.get('/api/admin/employees/'+str(bob['employee']['id'])+'/clients',headers=header(alice)).json()['companies'])==1
    assert app.patch('/api/admin/employees/'+str(bob['employee']['id'])+'/active',headers=header(alice),json={'active':False}).status_code==200
    assert app.get('/api/me',headers=header(bob)).status_code==401

def test_source_job_queue_is_per_employee_and_atomic(db):
    alice,bob=users(db);app=http(db)
    first=app.post('/api/search',headers=header(alice),json={'target':10})
    assert first.status_code==200
    assert app.post('/api/search',headers=header(alice),json={'target':10}).status_code==409
    assert app.post('/api/search',headers=header(bob),json={'target':10}).status_code==200
    assert app.post('/api/search/'+first.json()['id']+'/stop',headers=header(bob),json={}).status_code==404
    assert app.post('/api/search/'+first.json()['id']+'/stop',headers=header(alice),json={}).status_code==200
