from test_crm import db,pg,http,header,item
from test_customer_workspace import staff,setup_client

def populated(db):
    people=staff(db)
    ids=[setup_client(db,person,n+671) for n,person in enumerate(people)]
    return people,ids

def test_overview_is_default_leadership_scope_without_broadening_personal_clients(db):
    (a,b,m,d,outside),ids=populated(db);api=http(db)
    manager=api.get('/api/clients/overview',headers=header(m));assert manager.status_code==200,manager.text
    assert {c['id'] for c in manager.json()['companies']}==set(ids[:3])
    assert {c['id'] for c in api.get('/api/clients/overview',headers=header(d)).json()['companies']}==set(ids)
    assert api.get('/api/clients/overview',headers=header(a)).status_code==403
    assert api.get('/api/clients/overview',headers=header(b)).status_code==403
    assert api.get('/api/clients/overview').status_code==401
    assert {c['id'] for c in api.get('/api/clients',headers=header(m)).json()['companies']}=={ids[2]}
    for c in manager.json()['companies']:
        assert c['owner_first_name'] and c['owner_last_name'] and c['employee_id']
        assert c['stage']=='Согласование' and c['company_time']['utc_offset']==3
    assert not any(secret in manager.text for secret in ['password_hash','token_hash','mail_password'])

def test_owner_filter_and_selector_counts_respect_department_and_hierarchy(db):
    (a,b,m,d,outside),ids=populated(db);api=http(db)
    selected=api.get('/api/clients/overview',params={'employee_id':a['employee']['id']},headers=header(m))
    assert selected.status_code==200 and [c['id'] for c in selected.json()['companies']]==[ids[0]]
    assert {p['id'] for p in selected.json()['owners']}=={p['employee']['id'] for p in [a,b,m]}
    assert all(p['client_count']==1 for p in selected.json()['owners'])
    for person in [d,outside]:
        assert api.get('/api/clients/overview',params={'employee_id':person['employee']['id']},headers=header(m)).status_code==404
    assert api.get('/api/clients/overview',params={'employee_id':outside['employee']['id']},headers=header(d)).json()['companies'][0]['id']==ids[4]
    assert api.get('/api/clients/overview?order=not-a-sort',headers=header(d)).status_code==400
    assert api.get('/api/clients/overview?employee_id=not-a-uuid',headers=header(d)).status_code==422

def test_active_clients_only_and_personal_stage_catalog_is_preserved(db):
    (a,b,m,d,outside),ids=populated(db);api=http(db)
    with db.transaction(admin=True) as c:
        c.execute('UPDATE employee_clients SET active=false WHERE company_id=%s',(ids[1],))
    saved=api.put('/api/clients/stages',headers=header(m),json={'stages':['Согласование','Ждём договор']});assert saved.status_code==200
    data=api.get('/api/clients/overview',headers=header(m)).json()
    assert {c['id'] for c in data['companies']}=={ids[0],ids[2]}
    assert data['stages']==['Согласование']
    mine=api.get('/api/clients/overview',params={'employee_id':m['employee']['id']},headers=header(m)).json()
    assert mine['stages']==['Согласование','Ждём договор']
    assert {p['id'] for p in data['owners']}=={a['employee']['id'],m['employee']['id']}

def test_overview_keeps_owner_write_permissions_and_transfer_refreshes_owner(db):
    (a,b,m,d,outside),ids=populated(db);api=http(db)
    denied=api.patch('/api/clients/'+ids[0]+'/stage',headers=header(m),json={'stage':'Другой','revision':1})
    assert denied.status_code==404
    assert api.get('/api/companies/'+ids[0]+'/interactions',headers=header(m)).status_code==200
    moved=api.post('/api/companies/'+ids[0]+'/transfer',headers=header(m),json={'employee_id':b['employee']['id'],'revision':1,'note':'Из общего канбана'})
    assert moved.status_code==200,moved.text
    rows=api.get('/api/clients/overview',headers=header(m)).json()['companies']
    found=[c for c in rows if c['id']==ids[0]]
    assert len(found)==1 and found[0]['employee_id']==b['employee']['id'] and found[0]['owner_first_name']==b['employee']['first_name']
    assert not api.get('/api/clients/overview',params={'employee_id':a['employee']['id']},headers=header(m)).json()['companies']

def test_overview_excel_matches_owner_filter_and_never_exports_other_department(db):
    from io import BytesIO
    from openpyxl import load_workbook
    (a,b,m,d,outside),ids=populated(db);api=http(db)
    response=api.get('/api/clients/overview/export',headers=header(m));assert response.status_code==200
    rows=list(load_workbook(BytesIO(response.content),read_only=True).active.values)
    assert len(rows)==4 and rows[0][2]=='Ответственный'
    assert {r[2] for r in rows[1:]}=={p['employee']['first_name']+' '+p['employee']['last_name'] for p in [a,b,m]}
    selected=api.get('/api/clients/overview/export',params={'employee_id':a['employee']['id']},headers=header(m))
    assert len(list(load_workbook(BytesIO(selected.content),read_only=True).active.values))==2
    assert api.get('/api/clients/overview/export',headers=header(a)).status_code==403
