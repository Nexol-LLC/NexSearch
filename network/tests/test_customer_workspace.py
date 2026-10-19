from test_crm import db,pg,http,header,item
from test_v16 import staff as original_staff,task_body
from datetime import datetime,timezone
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
from io import BytesIO

def staff(db):
    people=original_staff(db)
    for person in people:person['employee']['id']=str(person['employee']['id'])
    return people

def setup_client(db,person,n=1):
    ident,_,_=db.ingest(item(n),person['employee']['id'])
    db.update_lead(person['employee']['id'],ident,'taken',True,'Первый контакт',0)
    db.update_stage(person['employee']['id'],ident,'Согласование',0)
    return str(ident)

def test_company_card_and_interactions_follow_department_permissions(db):
    a,b,manager,director,outside=staff(db);api=http(db);ident=setup_client(db,a)
    for p in [a,manager,director]:
        card=api.get('/api/companies/'+ident,headers=header(p));assert card.status_code==200,card.text
        assert card.json()['inn'] and card.json()['stage']=='Согласование'
        assert api.get('/api/companies/'+ident+'/interactions',headers=header(p)).json()['interactions']
    for p in [b,outside]:
        assert api.get('/api/companies/'+ident,headers=header(p)).status_code==404
        assert api.get('/api/companies/'+ident+'/interactions',headers=header(p)).status_code==404
    assert api.post('/api/companies/'+ident+'/interactions',headers=header(b),json={'description':'Чужая заметка'}).status_code==404
    assert api.post('/api/companies/'+ident+'/interactions',headers=header(manager),json={'description':'Обсудили условия перевозки'}).status_code==200
    db.audit(a['employee']['id'],'mail.send',ident,{'description':'Личное письмо, не общая история'})
    events=api.get('/api/companies/'+ident+'/interactions',headers=header(a)).json()['interactions']
    assert any(e['description']=='Обсудили условия перевозки' for e in events)
    assert not any(e['kind'].startswith('mail.') for e in events)

def test_transfer_is_atomic_preserves_stage_contacts_calls_and_moves_open_tasks(db):
    a,b,manager,director,outside=staff(db);api=http(db);ident=setup_client(db,a)
    contact=api.post('/api/contacts',headers=header(a),json={'company_id':ident,'name':'Контакт','phones':'+7 900 000 00 00'});assert contact.status_code==200
    opened=api.post('/api/tasks',headers=header(a),json=task_body(company_id=ident,title='Открытая задача')).json()
    closed=api.post('/api/tasks',headers=header(a),json=task_body(company_id=ident,title='Завершённая задача',status='done')).json()
    r=api.post('/api/companies/'+ident+'/transfer',headers=header(manager),json={'employee_id':b['employee']['id'],'revision':1,'note':'Продолжить согласование'});assert r.status_code==200,r.text
    assert r.json()['open_tasks']==1
    assert api.get('/api/companies/'+ident,headers=header(a)).status_code==404
    assert api.get('/api/companies/'+ident+'/interactions',headers=header(a)).status_code==404
    card=api.get('/api/companies/'+ident,headers=header(b)).json();assert card['stage']=='Согласование' and card['revision']==2
    assert api.get('/api/tasks/'+opened['id'],headers=header(b)).json()['assigned_by']==manager['employee']['id']
    assert api.get('/api/tasks/'+opened['id'],headers=header(a)).status_code==404
    assert api.get('/api/tasks/'+closed['id'],headers=header(a)).status_code==200
    assert api.get('/api/contacts?company_id='+ident,headers=header(b)).json()['contacts'][0]['name']=='Контакт'
    timeline=api.get('/api/companies/'+ident+'/interactions',headers=header(b)).json()['interactions']
    assert any(i['kind']=='call' and i['description']=='Первый контакт' for i in timeline)
    assert any(i['kind']=='company.transfer' for i in timeline)
    assert not api.get('/api/clients',headers=header(a)).json()['companies']
    assert len(api.get('/api/clients',headers=header(b)).json()['companies'])==1

def test_transfer_rejects_wrong_roles_other_department_stale_and_inactive(db):
    a,b,manager,director,outside=staff(db);api=http(db);ident=setup_client(db,a)
    url='/api/companies/'+ident+'/transfer';body={'employee_id':b['employee']['id'],'revision':1}
    assert api.post(url,headers=header(a),json=body).status_code==403
    assert api.post(url,headers=header(manager),json={**body,'employee_id':outside['employee']['id']}).status_code==404
    assert api.post(url,headers=header(manager),json={**body,'revision':0}).status_code==409
    with db.transaction(admin=True) as c:c.execute('UPDATE employees SET approved=false WHERE id=%s',(b['employee']['id'],))
    assert api.post(url,headers=header(director),json=body).status_code==404
    r=api.post(url,headers=header(director),json={**body,'employee_id':outside['employee']['id']});assert r.status_code==200,r.text
    assert api.get('/api/companies/'+ident,headers=header(manager)).status_code==404
    assert api.get('/api/companies/'+ident,headers=header(outside)).status_code==200

def test_two_simultaneous_transfers_have_one_winner(db):
    a,b,manager,director,outside=staff(db);api=http(db);ident=setup_client(db,a)
    def transfer(p):return api.post('/api/companies/'+ident+'/transfer',headers=header(p),json={'employee_id':b['employee']['id'],'revision':1}).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(transfer,[manager,director]))
    assert sorted(results)==[200,409]
    with db.transaction(admin=True) as c:
        assert c.execute('SELECT count(*) AS n FROM employee_clients WHERE company_id=%s',(ident,)).fetchone()['n']==1
        assert c.execute("SELECT count(*) AS n FROM audit_events WHERE kind='company.transfer'").fetchone()['n']==1

def test_task_catalog_distinguishes_found_clients_and_owner_scope(db):
    a,b,manager,director,outside=staff(db);api=http(db);client=setup_client(db,a)
    found,_,_=db.ingest(item(2),a['employee']['id'])
    rows=api.get('/api/tasks/companies',headers=header(a)).json()['companies']
    assert next(c for c in rows if c['id']==client)['is_client'] is True
    assert next(c for c in rows if c['id']==str(found))['is_client'] is False
    assert all(c['inn'] for c in rows)
    assert api.get('/api/tasks/companies?employee_id='+a['employee']['id'],headers=header(b)).status_code==403
    assert len(api.get('/api/tasks/companies?employee_id='+a['employee']['id'],headers=header(manager)).json()['companies'])==2

def test_fleet_export_has_tonnes_metres_and_no_location_or_total_weight(db):
    a,b,manager,director,outside=staff(db);api=http(db)
    body={'name':'Перевозчик','vehicles':[{'name':'Тент','kind':'tent','payload_kg':20500,'length_mm':13600,'width_mm':2450,'height_mm':2700,'weight_kg':8000,'location':'Казань'}]}
    assert api.post('/api/carriers/save',headers=header(a),json=body).status_code==200
    r=api.get('/api/export/carriers',headers=header(a));assert r.status_code==200
    fleet=load_workbook(BytesIO(r.content))['Автомобили'];headers=[cell.value for cell in fleet[1]]
    assert 'Где находится' not in headers and 'Вес машины, кг' not in headers
    assert fleet.cell(2,4).value==20.5 and fleet.cell(2,6).value==13.6
