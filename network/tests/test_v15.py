from test_crm import db,pg,users,item,http,header
from datetime import datetime,timezone,timedelta
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
import io,time,threading
from psycopg.types.json import Jsonb
from network.vat import calculate
from network.scheduling import enqueue_due
from network.search15 import SearchWorker

def test_vat_precise_add_extract_historical_and_without():
 assert calculate('100000','22')==dict(net='100000.00',vat='22000.00',gross='122000.00',rate='22',mode='add')
 assert calculate('122000','22','extract')['vat']=='22000.00'
 assert calculate('.03','22')['vat']=='0.01'
 assert calculate('120','20','extract')['net']=='100.00'
 assert calculate('200','0','add')['vat']=='0.00'
 assert calculate('200','22','without')['gross']=='200.00'

def test_count_boundaries_and_stopped_job_cannot_claim(db):
 a,b=users(db);api=http(db)
 for count in (0,51,None):assert api.post('/api/search',headers=header(a),json={'target':count}).status_code==422
 for n in range(20):db.ingest(item(n))
 j=api.post('/api/search',headers=header(a),json={'target':7}).json()
 assert api.post('/api/search/'+j['id']+'/stop',headers=header(b),json={}).status_code==404
 assert api.post('/api/search/'+j['id']+'/stop',headers=header(a),json={}).status_code==200
 assert db.claim_available(a['employee']['id'],7,j['id'])==[]
 assert not db.ingest(item(91),a['employee']['id'],job=j['id'])[2]
 assert not db.leads(a['employee']['id'])
 assert api.post('/api/search',headers=header(a),json={'target':50}).status_code==200

def test_cancel_during_verification_unblocks_and_prevents_assignment(db,monkeypatch):
 from network.cancellation import wait
 import network.search15 as module
 a,b=users(db);api=http(db);entered=threading.Event()
 source={'id':'egrul','name':'Тест','url':'https://example.org'};monkeypatch.setattr(module,'SOURCES',[source])
 class Checker:
  def key(self):return 'test'
  def enrich(self,value):entered.set();wait(30);return value
 worker=SearchWorker(db,Checker());api.app.state.worker=worker
 data=item(90);data['phones']=['+7 900 100 10 10']
 with db.transaction(admin=True) as c:c.execute('INSERT INTO scraper_candidates(url,source,payload) VALUES(%s,%s,%s)',(data['source_url'],'egrul',Jsonb(data)))
 j=api.post('/api/search',headers=header(a),json={'target':1}).json()
 with db.transaction(admin=True) as c:j=c.execute("UPDATE jobs SET status='running' WHERE id=%s RETURNING *",(j['id'],)).fetchone()
 worker.discover=lambda *args:iter(())
 with ThreadPoolExecutor(max_workers=1) as pool:
  future=pool.submit(worker.run,j);assert entered.wait(3)
  started=time.monotonic();response=api.post('/api/search/'+str(j['id'])+'/stop',headers=header(a),json={});future.result(timeout=3)
 assert response.status_code==200 and time.monotonic()-started<2
 assert worker.job(j['id'])['status']=='stopped' and not db.leads(a['employee']['id'])
 with db.transaction(admin=True) as c:assert c.execute("SELECT state FROM scraper_candidates").fetchone()['state']=='pending'

def test_schedules_are_private_claimed_once_and_defer_busy_employee(db):
 a,b=users(db);api=http(db)
 future=(datetime.now(timezone.utc)+timedelta(hours=1)).isoformat()
 schedule=api.post('/api/search/schedules',headers=header(a),json={'run_at':future,'target':8,'criteria':{}})
 assert schedule.status_code==200,schedule.text
 ident=schedule.json()['id']
 assert not api.get('/api/search/schedules',headers=header(b)).json()['schedules']
 assert api.delete('/api/search/schedules/'+ident,headers=header(b)).status_code==409
 job=api.post('/api/search',headers=header(a),json={'target':1}).json()
 with db.transaction(admin=True) as c:c.execute("UPDATE search_schedules SET run_at=now()-interval '1 minute'")
 assert enqueue_due(db)==0
 api.post('/api/search/'+job['id']+'/stop',headers=header(a),json={})
 with ThreadPoolExecutor(max_workers=3) as pool:assert sum(pool.map(lambda _:enqueue_due(db),range(3)))==1
 assert enqueue_due(db)==0
 own=api.get('/api/search/schedules',headers=header(a)).json()['schedules'][0]
 assert own['job_id'] and own['status']=='queued'

def test_carrier_atomic_fleet_contacts_terms_and_conflicts(db):
 a,b=users(db);api=http(db)
 fleet=[dict(name='Тент',kind='tent',length_mm=13600,width_mm=2450,height_mm=2700,payload_kg=20000,weight_kg=8000,volume_m3=90,location='Казань')]
 payload=dict(name='Перевозчик',phone='+7 900 111 22 33',email='transport@example.org',work_terms='Оплата через 5 дней',vat_mode='included',vat_rate=22,vehicles=fleet)
 response=api.post('/api/carriers/save',headers=header(a),json=payload);assert response.status_code==200,response.text
 carrier=response.json();shared=api.get('/api/carriers',headers=header(b)).json()['carriers'][0]
 assert shared['vehicles'][0]['volume_m3']==90 and shared['phone']==payload['phone'] and shared['work_terms']==payload['work_terms']
 updated=api.post('/api/carriers/save',headers=header(b),json=dict(payload,id=carrier['id'],revision=0,description='Общее описание'));assert updated.status_code==200
 assert api.post('/api/carriers/save',headers=header(a),json=dict(payload,id=carrier['id'],revision=0)).status_code==409
 bad=dict(payload,id=carrier['id'],revision=1,vehicles=[dict(fleet[0],payload_kg=-1)],phone='changed')
 assert api.post('/api/carriers/save',headers=header(a),json=bad).status_code==422
 assert api.get('/api/carriers',headers=header(a)).json()['carriers'][0]['phone']==payload['phone']
 legacy={'name':'Перевозчик','description':'Обновление из версии 1.4','revision':1}
 assert api.patch('/api/carriers/'+carrier['id'],headers=header(a),json=legacy).status_code==200
 oldfleet={k:v for k,v in fleet[0].items() if k not in ('weight_kg','volume_m3','location')}
 assert api.put('/api/carriers/'+carrier['id']+'/fleet',headers=header(b),json={'revision':2,'vehicles':[oldfleet]}).status_code==200
 preserved=api.get('/api/carriers',headers=header(a)).json()['carriers'][0]
 assert preserved['phone']==payload['phone'] and preserved['vehicles'][0]['location']=='Казань' and preserved['vehicles'][0]['volume_m3']==90

def test_lead_to_task_needs_no_client_transfer_and_is_private(db):
 a,b=users(db);api=http(db);ident,_,_=db.ingest(item(9),a['employee']['id'])
 payload=dict(company_id=str(ident),title='Позвонить из поиска',due_at=(datetime.now(timezone.utc)+timedelta(hours=1)).isoformat())
 assert api.post('/api/tasks',headers=header(a),json=payload).status_code==200
 assert api.post('/api/tasks',headers=header(b),json=payload).status_code==404
 assert len(api.get('/api/tasks',headers=header(a)).json()['tasks'])==1
 assert not api.get('/api/clients',headers=header(a)).json()['companies']

def test_orders_excel_typed_and_private(db):
 from test_v14 import order_payload
 from openpyxl import load_workbook
 from pathlib import Path
 a,b=users(db);api=http(db);ident,_,_=db.ingest(item(19),a['employee']['id'])
 payload=order_payload(str(ident));payload['title']='=2+2';payload['pickup_at']='2026-10-15T07:00:00Z'
 response=api.post('/api/orders',headers=header(a),json=payload);assert response.status_code==200,response.text
 response=api.get('/api/export/orders',headers=header(a));assert response.status_code==200
 book=load_workbook(io.BytesIO(response.content));assert book.sheetnames==['Заявки и рейсы','Грузы','Расходы']
 sheet=book.worksheets[0];assert sheet['B2'].value=='=2+2' and sheet['B2'].data_type=='s'
 assert sheet['D2'].data_type=='s' and sheet['L2'].data_type=='n'
 assert sheet['H2'].data_type=='d' and sheet['H2'].value.hour==10
 assert sheet.freeze_panes=='B2' and sheet.auto_filter.ref
 assert api.get('/api/export/orders?employee_id='+str(a['employee']['id']),headers=header(b)).status_code==403
 empty=load_workbook(io.BytesIO(api.get('/api/export/orders',headers=header(b)).content));assert empty.worksheets[0].max_row==1
 output=Path(__file__).parents[1]/'.runtime/qa15';output.mkdir(parents=True,exist_ok=True);(output/'orders.xlsx').write_bytes(response.content)

def test_four_employee_workers_parallel_disjoint_and_limited(db,monkeypatch):
 from network.auth import Authentication
 import network.search15 as module
 auth=Authentication(db);people=[auth.register('Логист',str(n),'password-parallel','parallel-'+str(n),'','127.0.0.1') for n in range(4)]
 with db.transaction(admin=True) as c:c.execute('UPDATE employees SET approved=true')
 source={'id':'egrul','name':'Тест','url':'https://example.org'};monkeypatch.setattr(module,'SOURCES',[source])
 class Checker:
  def key(self):return 'test'
  def enrich(self,value):time.sleep(.12);return value
 worker=SearchWorker(db,Checker());worker.discover=lambda *args:iter(())
 with db.transaction(admin=True) as c:
  for n in range(60):
   data=item(100+n);data['phones']=['+7 900 100 10 10']
   c.execute('INSERT INTO scraper_candidates(url,source,payload) VALUES(%s,%s,%s)',(data['source_url'],'egrul',Jsonb(data)))
  jobs=[c.execute("INSERT INTO jobs(id,employee_id,target,status,criteria) VALUES(%s,%s,10,'queued','{}') RETURNING *",(uuid4(),p['employee']['id'])).fetchone() for p in people]
 started=time.monotonic()
 worker.start()
 try:
  while time.monotonic()-started<8:
   if all(worker.job(j['id'])['status']=='completed' for j in jobs):break
   time.sleep(.05)
 finally:worker.stop()
 elapsed=time.monotonic()-started;sets=[{str(c['id']) for c in db.leads(p['employee']['id'])} for p in people]
 assert all(len(s)==10 for s in sets) and len(set.union(*sets))==40
 assert all(worker.job(j['id'])['added']==10 for j in jobs)
 # Artificial provider delay alone is >=4.8 s in a fully serial implementation.
 assert elapsed<3.5,elapsed
 print(f'BENCH15 four employees, 40 unique leads, mock check latency 120ms: {elapsed:.2f}s')

def test_checko_paces_starts_but_overlaps_network_waits(db,monkeypatch):
 from network.checko import Checko
 import network.checko as module
 checker=Checko(db);checker.key=lambda:'test';starts=[];lock=threading.Lock()
 class Response:
  status_code=200
  def json(self):return {'meta':{'status':'ok'},'data':{}}
 def get(*args,**kwargs):
  with lock:starts.append(time.monotonic())
  time.sleep(.6);return Response()
 monkeypatch.setattr(module.requests,'get',get)
 started=time.monotonic()
 with ThreadPoolExecutor(max_workers=4) as pool:list(pool.map(lambda _:checker.request('company','7700000000'),range(4)))
 elapsed=time.monotonic()-started
 assert elapsed<2.1 and all(b-a>=.34 for a,b in zip(starts,starts[1:]))

def test_checko_quota_message_and_circuit_breaker(db,monkeypatch):
 from network.checko import Checko,VerificationUnavailable
 import network.checko as module
 import pytest
 checker=Checko(db);checker.key=lambda:'test';calls=[]
 class Response:
  status_code=403
  def json(self):return {'meta':{'status':'error','message':'Превышен суточный лимит запросов для бесплатного тарифа','today_request_count':100,'balance':0}}
 monkeypatch.setattr(module.requests,'get',lambda *args,**kw:(calls.append(1) or Response()))
 for _ in range(3):
  with pytest.raises(VerificationUnavailable,match='суточный лимит'):checker.request('company','7700000000')
 assert len(calls)==1 and checker.metadata['today_request_count']==100

def test_live_host_scheduler_works_without_employee_ui(db):
 a,b=users(db);api=http(db);db.ingest(item(88))
 future=(datetime.now(timezone.utc)+timedelta(hours=1)).isoformat()
 schedule=api.post('/api/search/schedules',headers=header(a),json={'run_at':future,'target':1,'criteria':{}}).json()
 with db.transaction(admin=True) as c:c.execute("UPDATE search_schedules SET run_at=now()-interval '1 minute' WHERE id=%s",(schedule['id'],))
 class Checker:
  def key(self):return 'test'
 worker=SearchWorker(db,Checker());worker.start();started=time.monotonic();job=None
 try:
  while time.monotonic()-started<5:
   with db.transaction(admin=True) as c:row=c.execute('SELECT job_id FROM search_schedules WHERE id=%s',(schedule['id'],)).fetchone()
   if row['job_id']:
    job=worker.job(row['job_id'])
    if job['status']=='completed':break
   time.sleep(.05)
 finally:worker.stop()
 assert job and job['status']=='completed' and job['added']==1
